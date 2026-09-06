"""Сборка человекочитаемых Excel-выгрузок (.xlsx) по полигонам.

Раньше отчёт собирался на фронте в CSV из четырёх колонок — по нему нельзя
было понять ни откуда взялось значение NDVI, ни почему точка помечена
аномалией. Здесь выгружается весь ряд наблюдений с русскими заголовками и
единицами измерения (ровно те поля, что лежат в `NdviObservation`), плюс
отдельный лист с периодами аномалий и лист-сводка по участкам.

Файл собирается в память (`BytesIO`) — выгрузки маленькие (сотни строк на
полигон), поэтому временные файлы на диске не нужны.
"""

from __future__ import annotations

import io
import math
import re
import zipfile
from datetime import date as date_type
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.models.anomaly import AnomalyPeriod
from app.models.enums import NdviStatus
from app.models.polygon import Polygon
from app.models.timeseries import NdviObservation

# Подписи статусов дословно совпадают с фронтом
# (`frontend/lib/utils/ndvi_style.dart`, statusLabel) — иначе в приложении и
# в скачанном файле один и тот же участок назывался бы по-разному.
STATUS_LABELS: dict[NdviStatus, str] = {
    NdviStatus.normal: "Штатное развитие",
    NdviStatus.suppression: "Угнетение биомассы",
    NdviStatus.critical: "Критическая аномалия",
}

_HEADER_FILL = PatternFill("solid", fgColor="1F3B2C")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_DATE_FORMAT = "DD.MM.YYYY"
_NUM_FORMAT = "0.0000"

# (заголовок, атрибут модели, формат ячейки) — порядок колонок листа
# наблюдений. Держим одним списком, чтобы заголовок и значение нельзя было
# разъехать при правках.
_OBSERVATION_COLUMNS: list[tuple[str, str, str | None]] = [
    ("Дата", "date", _DATE_FORMAT),
    ("День года", "doy", "0"),
    ("NDVI (основной)", "primary_ndvi", _NUM_FORMAT),
    ("NDVI (восстановленный моделью)", "primary_ndvi_pred", _NUM_FORMAT),
    ("Пропуск восстановлен моделью", "is_synthetic_gap", None),
    ("NDVI Sentinel-2", "s2_ndvi", _NUM_FORMAT),
    ("EVI Sentinel-2", "s2_evi", _NUM_FORMAT),
    ("NDWI Sentinel-2", "s2_ndwi", _NUM_FORMAT),
    ("NDVI Landsat", "landsat_ndvi", _NUM_FORMAT),
    ("EVI Landsat", "landsat_evi", _NUM_FORMAT),
    ("NDWI Landsat", "landsat_ndwi", _NUM_FORMAT),
    ("NDVI MODIS", "modis_ndvi", _NUM_FORMAT),
    ("EVI MODIS", "modis_evi", _NUM_FORMAT),
    ("Температура, °C (ERA5)", "era5_temp_c", "0.00"),
    ("Осадки, мм (ERA5)", "era5_precip_mm", "0.00"),
    ("Норма NDVI (среднемноголетняя)", "ndvi_climatology_mean", _NUM_FORMAT),
    ("Стандартное отклонение нормы", "ndvi_climatology_std", _NUM_FORMAT),
    ("Z-оценка отклонения", "ndvi_zscore", "0.00"),
    ("Число лет в норме", "n_reference_years", "0"),
    ("Статус", "status", None),
    ("Культура", "crop_type", None),
]

_ANOMALY_HEADERS = [
    "Участок",
    "Начало периода",
    "Конец периода",
    "Длительность, дней",
    "Худшая Z-оценка",
    "Отклонение NDVI от нормы",
    "Статус",
    "Пояснение",
]

_SUMMARY_HEADERS = [
    "Участок",
    "Идентификатор",
    "Культура",
    "Площадь, га",
    "Наблюдений",
    "Из них восстановлено моделью",
    "Период с",
    "Период по",
    "Текущий статус",
    "Периодов аномалий",
]


def polygon_title(polygon: Polygon) -> str:
    """Человеческое имя участка: подпись, а если её нет — идентификатор."""
    return (polygon.label or "").strip() or polygon.id


def _area_hectares(points: list) -> float | None:
    """Площадь контура в гектарах. Тот же приём, что во фронте
    (`frontend/lib/utils/geo.dart`): формула шнурков в локальной
    равнопромежуточной проекции — для поля в несколько километров
    погрешность заведомо меньше точности самой отрисовки контура."""
    if not points or len(points) < 3:
        return None
    try:
        lats = [float(p[0]) for p in points]
        lons = [float(p[1]) for p in points]
    except (TypeError, ValueError, IndexError):
        return None

    lat0 = math.radians(sum(lats) / len(lats))
    r = 6371000.0
    xs = [math.radians(lon) * r * math.cos(lat0) for lon in lons]
    ys = [math.radians(lat) * r for lat in lats]

    area2 = 0.0
    for i in range(len(xs)):
        j = (i + 1) % len(xs)
        area2 += xs[i] * ys[j] - xs[j] * ys[i]
    return abs(area2) / 2 / 10_000


def _cell_value(obs: NdviObservation, attr: str):
    value = getattr(obs, attr, None)
    if attr == "is_synthetic_gap":
        return "Да" if value else "Нет"
    if attr == "status":
        return STATUS_LABELS.get(value, "") if value is not None else "Нет данных"
    if attr == "crop_type":
        return value or "—"
    if isinstance(value, float):
        # Пишем уже округлённое значение, а не только формат ячейки: индексы
        # ДЗЗ дальше 4-го знака не имеют смысла, а «сырые» 0.14248919525321
        # в строке формул — ровно то, из-за чего выгрузка выглядела нечитаемой.
        return round(value, 4)
    return value


def _style_header(ws: Worksheet, headers: list[str]) -> None:
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    # Заголовок не должен уезжать при прокрутке — ряд длинный, иначе на
    # середине файла уже не понять, что за колонка.
    ws.freeze_panes = "A2"


def _autofit(ws: Worksheet, headers: list[str], max_width: int = 60) -> None:
    """Ширина по факту содержимого: openpyxl не умеет «автоподбор», как
    Excel, поэтому считаем самую длинную строку в колонке сами."""
    for col, title in enumerate(headers, start=1):
        longest = len(str(title))
        for row in range(2, ws.max_row + 1):
            value = ws.cell(row=row, column=col).value
            if value is None:
                continue
            text = value.strftime("%d.%m.%Y") if isinstance(value, (datetime, date_type)) else str(value)
            longest = max(longest, len(text))
        ws.column_dimensions[get_column_letter(col)].width = min(max_width, max(10, longest + 2))


def _sheet_title(base: str, used: set[str]) -> str:
    """Имя листа Excel: не длиннее 31 символа, без `[]:*?/\\` и уникальное
    в книге — иначе openpyxl молча переименует лист или упадёт."""
    cleaned = re.sub(r"[\[\]:*?/\\]", " ", base).strip() or "Участок"
    title = cleaned[:31]
    suffix = 2
    while title.lower() in used:
        tail = f" ({suffix})"
        title = cleaned[: 31 - len(tail)] + tail
        suffix += 1
    used.add(title.lower())
    return title


def _write_observations(ws: Worksheet, observations: list[NdviObservation]) -> None:
    headers = [title for title, _, _ in _OBSERVATION_COLUMNS]
    _style_header(ws, headers)
    for row_idx, obs in enumerate(sorted(observations, key=lambda o: o.date), start=2):
        for col_idx, (_, attr, fmt) in enumerate(_OBSERVATION_COLUMNS, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=_cell_value(obs, attr))
            if fmt is not None and cell.value is not None:
                cell.number_format = fmt
    _autofit(ws, headers)


def _write_anomalies(ws: Worksheet, rows: list[tuple[str, AnomalyPeriod]]) -> None:
    _style_header(ws, _ANOMALY_HEADERS)
    for row_idx, (title, anomaly) in enumerate(rows, start=2):
        duration = (anomaly.end_date - anomaly.start_date).days + 1
        values = [
            title,
            anomaly.start_date,
            anomaly.end_date,
            duration,
            round(anomaly.min_z_score, 4),
            round(anomaly.deviation, 4),
            STATUS_LABELS.get(anomaly.severity, str(anomaly.severity)),
            anomaly.explanation,
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if col_idx in (2, 3):
                cell.number_format = _DATE_FORMAT
            elif col_idx in (5, 6):
                cell.number_format = "0.00"
        ws.cell(row=row_idx, column=8).alignment = Alignment(wrap_text=True, vertical="top")
    if not rows:
        ws.cell(row=2, column=1, value="Аномальных периодов не зафиксировано")
    _autofit(ws, _ANOMALY_HEADERS)


def _write_summary(ws: Worksheet, bundles: list["PolygonExport"]) -> None:
    _style_header(ws, _SUMMARY_HEADERS)
    for row_idx, bundle in enumerate(bundles, start=2):
        observations = sorted(bundle.observations, key=lambda o: o.date)
        area = _area_hectares(bundle.polygon.points)
        last_status = observations[-1].status if observations else None
        values = [
            polygon_title(bundle.polygon),
            bundle.polygon.id,
            bundle.polygon.crop_type or "—",
            round(area, 1) if area is not None else None,
            len(observations),
            sum(1 for o in observations if o.is_synthetic_gap),
            observations[0].date if observations else None,
            observations[-1].date if observations else None,
            STATUS_LABELS.get(last_status, "Нет данных") if last_status else "Нет данных",
            len(bundle.anomalies),
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if col_idx == 4 and value is not None:
                cell.number_format = "0.0"
            elif col_idx in (7, 8) and value is not None:
                cell.number_format = _DATE_FORMAT
    _autofit(ws, _SUMMARY_HEADERS)


class PolygonExport:
    """Всё, что нужно для выгрузки одного участка: сам полигон, его ряд
    наблюдений и периоды аномалий."""

    __slots__ = ("polygon", "observations", "anomalies")

    def __init__(
        self,
        polygon: Polygon,
        observations: list[NdviObservation],
        anomalies: list[AnomalyPeriod],
    ) -> None:
        self.polygon = polygon
        self.observations = observations
        self.anomalies = anomalies


def _workbook_bytes(wb: Workbook) -> bytes:
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_polygon_workbook(bundle: PolygonExport) -> bytes:
    """Книга по одному участку: «Сводка» + «Наблюдения» + «Аномалии»."""
    wb = Workbook()
    summary = wb.active
    summary.title = "Сводка"
    _write_summary(summary, [bundle])
    _write_observations(wb.create_sheet("Наблюдения"), bundle.observations)
    _write_anomalies(
        wb.create_sheet("Аномалии"),
        [(polygon_title(bundle.polygon), a) for a in sorted(bundle.anomalies, key=lambda a: a.start_date)],
    )
    return _workbook_bytes(wb)


def build_combined_workbook(bundles: list[PolygonExport]) -> bytes:
    """Одна книга на несколько участков: сводка, лист наблюдений на каждый
    участок и общий лист аномалий."""
    wb = Workbook()
    summary = wb.active
    summary.title = "Сводка"
    _write_summary(summary, bundles)

    used_titles = {"сводка", "аномалии"}
    for bundle in bundles:
        ws = wb.create_sheet(_sheet_title(polygon_title(bundle.polygon), used_titles))
        _write_observations(ws, bundle.observations)

    all_anomalies = [
        (polygon_title(b.polygon), a)
        for b in bundles
        for a in sorted(b.anomalies, key=lambda a: a.start_date)
    ]
    _write_anomalies(wb.create_sheet("Аномалии"), all_anomalies)
    return _workbook_bytes(wb)


def file_stem(polygon: Polygon) -> str:
    """Имя файла: подпись участка латиницей/кириллицей без спецсимволов —
    браузеры и архиваторы плохо переносят `/`, кавычки и переводы строк."""
    base = re.sub(r"[^\w\-. ]+", "_", polygon_title(polygon), flags=re.UNICODE).strip()
    return (base or polygon.id)[:60]


def build_zip(bundles: list[PolygonExport]) -> bytes:
    """Zip из отдельных .xlsx на каждый участок — пользователь просил
    именно так, когда участков несколько (по файлу на участок читать
    проще, чем искать нужный лист в общей книге)."""
    buffer = io.BytesIO()
    used_names: set[str] = set()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for bundle in bundles:
            name = f"{file_stem(bundle.polygon)}.xlsx"
            suffix = 2
            while name.lower() in used_names:
                name = f"{file_stem(bundle.polygon)} ({suffix}).xlsx"
                suffix += 1
            used_names.add(name.lower())
            archive.writestr(name, build_polygon_workbook(bundle))
        # Общая книга-сводка рядом с пофайловыми выгрузками — чтобы сравнить
        # участки между собой, не открывая каждый файл отдельно.
        archive.writestr("Сводка по участкам.xlsx", build_combined_workbook(bundles))
    return buffer.getvalue()
